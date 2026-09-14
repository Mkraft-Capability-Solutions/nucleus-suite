import openpyxl
import json
import re

wb = openpyxl.load_workbook('documentation/sheets/Nucleus_Forms_and_Fields_Complete_MKraft.xlsx', data_only=True)
field_sheet = wb['02_Field_Spec']
picklist_sheet = wb['03_Picklists']

# 1. Parse Picklists
picklists = {}
for r in list(picklist_sheet.iter_rows(values_only=True))[3:]:
    code = str(r[0] or '').strip()
    name = str(r[1] or '').strip()
    raw_vals = str(r[4] or '').strip()
    if code and raw_vals:
        vals = [v.strip() for v in raw_vals.split('·') if v.strip()]
        picklists[code] = {
            'code': code,
            'name': name,
            'values': vals,
            'options': [{'value': v.lower().replace(' ', '_').replace('/', '_'), 'label': v} for v in vals]
        }

# 2. Parse 02_Field_Spec
forms = {}
current_form = None

rows = list(field_sheet.iter_rows(values_only=True))

for r in rows[3:]:
    col0 = str(r[0] or '').strip()
    if col0.startswith('§'):
        parts = [p.strip() for p in col0.split('·')]
        form_id = parts[0].replace('§', '').strip()
        form_name = parts[1] if len(parts) > 1 else ''
        module = parts[2] if len(parts) > 2 else ''
        scr_id = parts[3] if len(parts) > 3 else ''
        field_count_str = parts[4] if len(parts) > 4 else ''
        current_form = form_id
        forms[current_form] = {
            'form_id': form_id,
            'form_name': form_name,
            'module': module,
            'screen_id': scr_id,
            'fields': []
        }
    elif col0.startswith('FRM-'):
        current_form = col0
        if current_form not in forms:
            forms[current_form] = {
                'form_id': col0,
                'form_name': str(r[1] or '').strip(),
                'module': str(r[2] or '').strip(),
                'screen_id': '',
                'fields': []
            }
        
        field_name = str(r[5] or '').strip()
        field_label = str(r[4] or '').strip()
        control = str(r[6] or '').strip()
        data_type = str(r[7] or '').strip()
        mandatory = True if str(r[8] or '').strip().upper() == 'Y' else False
        default_val = str(r[9] or '').strip()
        validation = str(r[10] or '').strip()
        pl_code = str(r[11] or '').strip()
        section = str(r[3] or '').strip()
        visibility = str(r[12] or '').strip()
        workday = str(r[13] or '').strip()
        note = str(r[14] or '').strip()

        # Derive camelCase key
        clean_name = re.sub(r'[^a-zA-Z0-9_]', '_', field_name)
        parts = clean_name.split('_')
        camel_key = parts[0].lower() + ''.join(p.capitalize() for p in parts[1:] if p)
        if not camel_key:
            # Fallback from label
            clean_label = re.sub(r'[^a-zA-Z0-9_ ]', '', field_label)
            lparts = clean_label.split()
            camel_key = lparts[0].lower() + ''.join(p.capitalize() for p in lparts[1:]) if lparts else 'field'

        # Map control type
        c_low = control.lower()
        d_low = data_type.lower()
        if 'time' in c_low or 'time' in d_low:
            ftype = 'time'
        elif 'date' in c_low or 'date' in d_low:
            ftype = 'date'
        elif 'toggle' in c_low or 'checkbox' in c_low or 'bool' in d_low:
            ftype = 'checkbox'
        elif 'area' in c_low or 'multiline' in c_low or 'rich' in c_low:
            ftype = 'textarea'
        elif 'number' in c_low or 'int' in d_low or 'dec' in d_low or 'num' in d_low or 'currency' in c_low:
            ftype = 'number'
        elif 'select' in c_low or 'lookup' in c_low or 'dropdown' in c_low or (pl_code and pl_code != '—'):
            ftype = 'select'
        elif 'file' in c_low or 'upload' in c_low or 'document' in c_low or 'attachment' in c_low:
            ftype = 'file'
        else:
            ftype = 'text'

        options = []
        if pl_code and pl_code in picklists:
            options = picklists[pl_code]['options']
        elif ftype == 'select' and validation and '|' in validation:
            opts = [v.strip() for v in validation.split('|') if v.strip()]
            options = [{'value': v.lower().replace(' ', '_'), 'label': v} for v in opts]

        forms[current_form]['fields'].append({
            'key': camel_key,
            'label': field_label,
            'required': mandatory,
            'type': ftype,
            'control': control,
            'dataType': data_type,
            'section': section,
            'default': default_val if default_val != '—' else None,
            'validation': validation if validation != '—' else None,
            'picklist': pl_code if pl_code != '—' else None,
            'workday': workday if workday != '—' else None,
            'note': note if note != '—' else None,
            'options': options
        })

print(f"Parsed {len(forms)} forms with complete fields.")
with open('documentation/sheets/parsed_forms_spec.json', 'w') as out:
    json.dump(forms, out, indent=2)
