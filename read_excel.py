import sys
try:
    import openpyxl
except ImportError:
    print('openpyxl is not installed. Please install it.')
    sys.exit(1)

files = {
    'forms': 'docs/sheets/Nucleus_Forms_and_Fields_Complete_MKraft.xlsx',
    'process': 'docs/sheets/Nucleus_Process_Flows_and_Process_Maps_v1.0.xlsx',
    'demo': 'docs/sheets/HR Demo Points.xlsx',
}

for key, path in files.items():
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        print(f'=== {key}: {path} ===')
        print(f'Sheets: {wb.sheetnames}')
        for name in wb.sheetnames:
            ws = wb[name]
            print(f'  Sheet [{name}]: {ws.max_row} rows x {ws.max_column} cols')
        wb.close()
    except Exception as e:
        print(f'ERROR {key}: {e}')
