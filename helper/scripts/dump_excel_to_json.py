import openpyxl
import json
import os
import datetime

file_path = r"C:\Users\singh\Downloads\Nucleus_HRMS_Demo_Dataset_v1.0.xlsx"
out_dir = r"c:\Users\singh\Downloads\mkraft\mkraft-hrms\scripts\excel_data"
os.makedirs(out_dir, exist_ok=True)

wb = openpyxl.load_workbook(file_path, data_only=True)

def serialize_val(val):
    if val is None:
        return None
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.isoformat()
    if isinstance(val, (int, float)):
        return val
    return str(val).strip()

manifest = {}

for sheet_name in wb.sheetnames:
    sheet = wb[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        continue
    headers = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(rows[0])]
    data = []
    for r in rows[1:]:
        # check if row is empty
        if not any(c is not None for c in r):
            continue
        row_obj = {}
        for h, val in zip(headers, r):
            row_obj[h] = serialize_val(val)
        data.append(row_obj)
    
    clean_name = sheet_name.replace(" ", "_").lower()
    json_path = os.path.join(out_dir, f"{clean_name}.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    
    manifest[sheet_name] = {
        "file": f"{clean_name}.json",
        "count": len(data),
        "headers": headers
    }

manifest_path = os.path.join(out_dir, "manifest.json")
with open(manifest_path, "w", encoding="utf-8") as f:
    json.dump(manifest, f, indent=2)

print(f"Extracted {len(manifest)} sheets to {out_dir}")
for k, v in manifest.items():
    print(f"  {k}: {v['count']} records")
