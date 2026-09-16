import openpyxl
import sys
import json

file_path = r"C:\Users\singh\Downloads\Nucleus_HRMS_Demo_Dataset_v1.0.xlsx"

try:
    wb = openpyxl.load_workbook(file_path, data_only=True)
except Exception as e:
    print(f"Error loading workbook: {e}")
    sys.exit(1)

print("=== SHEET NAMES ===")
sheet_names = wb.sheetnames
for idx, name in enumerate(sheet_names, 1):
    print(f"{idx}. {name}")

print("\n=== README TAB CONTENT ===")
readme_sheet = None
for name in sheet_names:
    if "readme" in name.lower():
        readme_sheet = wb[name]
        break

if readme_sheet:
    for row in readme_sheet.iter_rows(values_only=True):
        # Filter out empty lines
        non_empty = [str(cell) for cell in row if cell is not None and str(cell).strip()]
        if non_empty:
            print(" | ".join(non_empty))
else:
    print("No sheet with 'readme' in its name found!")

print("\n=== SHEET OVERVIEWS ===")
overview = {}
for name in sheet_names:
    sheet = wb[name]
    max_r = sheet.max_row
    max_c = sheet.max_column
    headers = []
    if max_r and max_r > 0:
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])
        headers = [str(h) for h in first_row if h is not None]
    overview[name] = {
        "rows": max_r,
        "cols": max_c,
        "headers": headers[:15]  # first 15 headers
    }

print(json.dumps(overview, indent=2))
