import openpyxl

file_path = r"C:\Users\singh\Downloads\Nucleus_HRMS_Demo_Dataset_v1.0.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

readme = None
for name in wb.sheetnames:
    if "readme" in name.lower() or "00" in name:
        readme = wb[name]
        print(f"--- SHEET: {name} ---")
        for row in readme.iter_rows(values_only=True):
            vals = [str(c) for c in row if c is not None and str(c).strip()]
            if vals:
                print(" | ".join(vals))

print("\n--- ALL SHEET NAMES ---")
for i, s in enumerate(wb.sheetnames):
    print(f"{i}: {s}")
