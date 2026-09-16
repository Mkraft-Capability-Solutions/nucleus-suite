import openpyxl

file_path = r"C:\Users\singh\Downloads\Nucleus_HRMS_Demo_Dataset_v1.0.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

sheets_to_inspect = [
    "14_Leave_Types", "15_Leave_Accrual_Policy", "16_Leave_Proration", "17_Leave_Ledger",
    "18_Leave_Requests", "19_CompOff_Ledger", "20_Punch_Events", "21_Expected_Attendance",
    "22_Gate_Pass", "23_Overtime_Register", "24_Loans", "25_Loan_Guarantors",
    "26_Payroll_Runs", "27_Exit_Clearance", "28_Statutory_Calendar", "29_Statutory_Forms",
    "30_Recognition_Referral", "31_Announcements", "32_Letter_Templates", "33_Assets",
    "34_Induction_Learning", "35_GL_Mapping", "36_ERP_Inbound_Master", "37_Requisitions"
]

for s_name in sheets_to_inspect:
    sheet = wb[s_name]
    print(f"\n==================== {s_name} (rows: {sheet.max_row}) ====================")
    rows = list(sheet.iter_rows(values_only=True))
    if rows:
        headers = [str(c) if c is not None else "" for c in rows[0]]
        print("HEADERS:", headers)
        for r in rows[1:3]: # sample 2 rows
            vals = [str(c) if c is not None else "" for c in r]
            print("ROW:", vals)
