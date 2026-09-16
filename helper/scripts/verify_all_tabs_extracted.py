import openpyxl
import os
import json

file_path = r"C:\Users\singh\Downloads\Nucleus_HRMS_Demo_Dataset_v1.0.xlsx"
out_dir = r"c:\Users\singh\Downloads\mkraft\mkraft-hrms\scripts\excel_data"

wb = openpyxl.load_workbook(file_path, data_only=True)

print(f"Total sheets found in workbook: {len(wb.sheetnames)}")
print("List of all sheets and visibility:")
for i, name in enumerate(wb.sheetnames):
    ws = wb[name]
    visibility = ws.sheet_state
    print(f"[{i:02d}] {name} (state={visibility}, max_row={ws.max_row}, max_col={ws.max_column})")

# Check what was saved to JSON
json_files = os.listdir(out_dir)
print(f"\nTotal JSON files extracted in {out_dir}: {len(json_files)}")

with open(os.path.join(out_dir, "manifest.json"), "r") as f:
    manifest = json.load(f)

for name, meta in manifest.items():
    file_name = meta["file"]
    file_path_json = os.path.join(out_dir, file_name)
    with open(file_path_json, "r", encoding="utf-8") as jf:
        data = json.load(jf)
    print(f"Sheet '{name}': {len(data)} rows extracted, {len(meta['headers'])} headers")
