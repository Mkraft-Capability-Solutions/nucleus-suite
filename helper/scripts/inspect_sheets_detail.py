import openpyxl
import json

file_path = r"C:\Users\singh\Downloads\Nucleus_HRMS_Demo_Dataset_v1.0.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

sheets_to_inspect = [
    "02_Legal_Entities", "03_Locations", "04_Org_Units", "05_Designations",
    "06_Worker_Classes", "07_Shifts", "08_Attendance_Rules", "09_Holiday_Calendar",
    "10_Sanctioned_Manpower", "11_Positions", "12_Employees", "13_Salary_Structure"
]

for s_name in sheets_to_inspect:
    sheet = wb[s_name]
    print(f"\n==================== {s_name} (rows: {sheet.max_row}) ====================")
    rows = list(sheet.iter_rows(values_only=True))
    if rows:
        headers = [str(c) if c is not None else "" for c in rows[0]]
        print("HEADERS:", headers)
        for r in rows[1:4]: # sample 3 rows
            vals = [str(c) if c is not None else "" for c in r]
            print("ROW:", vals)
